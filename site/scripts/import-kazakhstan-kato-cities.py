#!/usr/bin/env python3
"""Build Kazakhstan's current statutory city catalog from KATO.

Usage:
  py -3.12 scripts/import-kazakhstan-kato-cities.py \
    <KATO_17.07.2026.xlsx> <KZ.txt> [output.csv]

The KATO snapshot contains every administrative object. This importer keeps
only rows whose Russian name is prefixed ``г.`` (city), excluding ``г.а.``
(city administration) and all settlements below city level.
"""

from __future__ import annotations

import csv
import hashlib
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError as error:
    raise SystemExit("缺少 openpyxl；请使用项目工作区自带的 Python 依赖运行") from error


EXPECTED_KATO_HASH = (
    "5346ADE0EC37A5A6F6B186E626527A17685121EB73B57152AA09E7F370FCFF09"
)
EXPECTED_GEONAMES_HASH = (
    "15F639EB644F3AB62312DB175B84091EC8A6433324EC8E3E90E0E4C0EBAB01EB"
)

SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_OUTPUT = (
    SCRIPT_DIR.parent / "data" / "sources" / "kz-kato-cities-2026-07-17.csv"
)

REGIONS = {
    "10": ("Abay Region", "12510143"),
    "11": ("Akmola Region", "03"),
    "15": ("Aktobe Region", "04"),
    "19": ("Almaty Region", "01"),
    "23": ("Atyrau Region", "06"),
    "27": ("West Kazakhstan Region", "07"),
    "31": ("Jambyl Region", "17"),
    "33": ("Jetisu Region", "12510144"),
    "35": ("Karaganda Region", "12"),
    "39": ("Kostanay Region", "13"),
    "43": ("Kyzylorda Region", "14"),
    "47": ("Mangystau Region", "09"),
    "55": ("Pavlodar Region", "11"),
    "59": ("North Kazakhstan Region", "16"),
    "61": ("Turkistan Region", "10"),
    "62": ("Ulytau Region", "12510145"),
    "63": ("East Kazakhstan Region", "15"),
    "71": ("Astana", "05"),
    "75": ("Almaty", "02"),
    "79": ("Shymkent", "1537272"),
}

EXPECTED_LEVEL_COUNTS = {
    "republican_significance_city": 3,
    "regional_significance_city": 39,
    "district_significance_city": 48,
}

# Current city names that differ from the corresponding GeoNames record, or
# whose country snapshot has a misleading same-name point. Alatau was created
# around the former Zhetigen settlement in 2024. Baikonur's city point is in
# GeoNames admin1 08 because of its special leased-territory status.
POINT_OVERRIDES = {
    "111610000": ("Qosshy", "1522202"),
    "191810000": ("Alatau", "1520706"),
    "334820100": ("Sarqan", "1519691"),
    "431910000": ("Baikonur", "1521368"),
}

# Zhem has no populated-place record in the pinned GeoNames snapshot. The
# point is the current city coordinate published by the Aktobe regional land
# authority; the GeoNames id intentionally remains empty.
MANUAL_POINTS = {
    "154825100": ("Zhem", "58.071269", "48.775253"),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def require_hash(path: Path, expected: str) -> None:
    actual = sha256(path)
    if actual != expected:
        raise SystemExit(f"{path} 快照哈希变化：{actual}")


def normalized_name(value: str) -> str:
    value = unicodedata.normalize("NFKD", value).lower().strip()
    value = re.sub(r"^(?:г\.|қ\.)", "", value)
    value = re.sub(r"[қќ]\.$", "", value).strip()
    return "".join(character for character in value if character.isalnum())


def parse_kato(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "katonew1" not in workbook.sheetnames:
        raise SystemExit("KATO 工作簿缺少 katonew1 工作表")
    worksheet = workbook["katonew1"]
    rows = worksheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    expected_headers = ["te", "ab", "cd", "ef", "hij", "k", "kaz_name", "rus_name", "nn"]
    if headers != expected_headers:
        raise SystemExit(f"KATO 表头变化：{headers}")

    cities = []
    for values in rows:
        row = dict(zip(headers, values))
        russian_name = str(row["rus_name"] or "").strip()
        if not russian_name.startswith("г.") or russian_name.startswith("г.а."):
            continue
        code = str(row["te"] or "").strip()
        if not re.fullmatch(r"\d{9}", code):
            raise SystemExit(f"KATO 城市代码格式异常：{code}")
        region_code = code[:2]
        if region_code not in REGIONS:
            raise SystemExit(f"KATO 城市所属地区未知：{code}")
        ef = str(row["ef"] or "").zfill(2)
        hij = str(row["hij"] or "").zfill(3)
        if region_code in {"71", "75", "79"}:
            level = "republican_significance_city"
        elif ef == "10" and hij == "000":
            level = "regional_significance_city"
        else:
            level = "district_significance_city"
        cities.append(
            {
                "code": code,
                "kazakh_name": str(row["kaz_name"] or "").strip(),
                "russian_name": russian_name,
                "region_code": region_code,
                "city_level": level,
            }
        )

    counts = Counter(city["city_level"] for city in cities)
    if len(cities) != 90 or counts != EXPECTED_LEVEL_COUNTS:
        raise SystemExit(
            f"KATO 城市分类变化：共 {len(cities)} 个，{dict(sorted(counts.items()))}"
        )
    if len({city["code"] for city in cities}) != len(cities):
        raise SystemExit("KATO 城市稳定 ID 重复")
    return cities


def parse_geonames(path: Path) -> tuple[dict[str, dict[str, object]], dict[tuple[str, str], list[dict[str, object]]]]:
    places_by_id: dict[str, dict[str, object]] = {}
    places_by_admin_and_name: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    with path.open(encoding="utf-8") as source:
        for line in source:
            values = line.rstrip("\n").split("\t")
            if len(values) < 19 or values[6] != "P":
                continue
            names = {values[1], values[2], *values[3].split(",")}
            place: dict[str, object] = {
                "id": values[0],
                "name": values[1],
                "ascii_name": values[2],
                "latitude": values[4],
                "longitude": values[5],
                "feature_code": values[7],
                "admin1": values[10],
                "population": int(values[14] or 0),
            }
            places_by_id[values[0]] = place
            for name in names:
                key = normalized_name(name)
                if key:
                    places_by_admin_and_name[(values[10], key)].append(place)
    return places_by_id, places_by_admin_and_name


def point_rank(place: dict[str, object]) -> tuple[int, int, int]:
    important = int(str(place["feature_code"]) in {"PPLC", "PPLA", "PPLA2", "PPLA3"})
    return int(place["population"]), important, int(place["id"])


def choose_point(
    city: dict[str, str],
    places_by_id: dict[str, dict[str, object]],
    places_by_admin_and_name: dict[tuple[str, str], list[dict[str, object]]],
) -> tuple[str, str, str, str]:
    manual = MANUAL_POINTS.get(city["code"])
    if manual:
        name, longitude, latitude = manual
        return name, "", longitude, latitude

    override = POINT_OVERRIDES.get(city["code"])
    if override:
        name, point_id = override
        point = places_by_id.get(point_id)
        if point is None:
            raise SystemExit(f"固定中心点不存在：{city['code']} -> {point_id}")
        return name, point_id, str(point["longitude"]), str(point["latitude"])

    _, expected_admin1 = REGIONS[city["region_code"]]
    candidates: dict[str, dict[str, object]] = {}
    for official_name in (city["kazakh_name"], city["russian_name"]):
        key = (expected_admin1, normalized_name(official_name))
        for place in places_by_admin_and_name.get(key, []):
            candidates[str(place["id"])] = place
    if not candidates:
        raise SystemExit(
            f"城市没有行政区内精确中心点：{city['code']} / {city['kazakh_name']}"
        )
    point = max(candidates.values(), key=point_rank)
    name = str(point["ascii_name"] or point["name"])
    return name, str(point["id"]), str(point["longitude"]), str(point["latitude"])


def build_rows(
    cities: list[dict[str, str]],
    places_by_id: dict[str, dict[str, object]],
    places_by_admin_and_name: dict[tuple[str, str], list[dict[str, object]]],
) -> list[dict[str, str]]:
    output = []
    for city in cities:
        name, point_id, longitude_text, latitude_text = choose_point(
            city, places_by_id, places_by_admin_and_name
        )
        longitude = float(longitude_text)
        latitude = float(latitude_text)
        if not (46 <= longitude <= 88 and 40 <= latitude <= 56):
            raise SystemExit(f"{city['code']} / {name} 中心点超出哈萨克斯坦范围")
        output.append(
            {
                "administrative_code": city["code"],
                "name": name,
                "admin_area": REGIONS[city["region_code"]][0],
                "city_level": city["city_level"],
                "geonames_id": point_id,
                "longitude": f"{longitude:.6f}",
                "latitude": f"{latitude:.6f}",
            }
        )

    for field, values in (
        ("稳定 ID", [row["administrative_code"] for row in output]),
        ("GeoNames ID", [row["geonames_id"] for row in output if row["geonames_id"]]),
        ("中心坐标", [(row["longitude"], row["latitude"]) for row in output]),
    ):
        duplicates = [value for value, count in Counter(values).items() if count > 1]
        if duplicates:
            raise SystemExit(f"生成结果存在重复{field}：{duplicates}")
    if len(output) != 90:
        raise SystemExit(f"生成城市数应为 90，实际为 {len(output)}")
    return sorted(
        output,
        key=lambda row: (row["admin_area"], row["name"], row["administrative_code"]),
    )


def write_csv(rows: list[dict[str, str]], path: Path) -> None:
    headers = [
        "administrative_code",
        "name",
        "admin_area",
        "city_level",
        "geonames_id",
        "longitude",
        "latitude",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    if len(sys.argv) not in {3, 4}:
        raise SystemExit(
            "用法：py -3.12 scripts/import-kazakhstan-kato-cities.py "
            "<KATO_17.07.2026.xlsx> <KZ.txt> [output.csv]"
        )
    kato_path = Path(sys.argv[1]).resolve()
    geonames_path = Path(sys.argv[2]).resolve()
    output_path = Path(sys.argv[3]).resolve() if len(sys.argv) == 4 else DEFAULT_OUTPUT
    require_hash(kato_path, EXPECTED_KATO_HASH)
    require_hash(geonames_path, EXPECTED_GEONAMES_HASH)
    places_by_id, places_by_admin_and_name = parse_geonames(geonames_path)
    rows = build_rows(
        parse_kato(kato_path), places_by_id, places_by_admin_and_name
    )
    write_csv(rows, output_path)
    print("已生成 90 个哈萨克斯坦法定城市中心点（3/39/48）")


if __name__ == "__main__":
    main()
