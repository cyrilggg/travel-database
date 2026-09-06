#!/usr/bin/env python3
"""Build Georgia's current statutory city catalog.

Usage:
  py -3.12 scripts/import-georgia-statutory-cities.py \
    <population-admin-units-2024.xlsx> <GE.txt> [output.csv]
"""

from __future__ import annotations

import csv
import hashlib
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_POPULATION_HASH = (
    "D6348F94CB4827C4C9433F3EE5820EC35DC016254316AE65EB07FE923EDC0DE1"
)
EXPECTED_GEONAMES_HASH = (
    "C931624C25587EEECA8ABB6490C572122C85188B5DDD765F5B1589E967F816B7"
)

SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_OUTPUT = (
    SCRIPT_DIR.parent / "data" / "sources" / "ge-statutory-cities-2026-09-06.csv"
)

# Official English city name, containing administrative unit, statutory level.
# The first 55 names are read back from Geostat's final 2024 census workbook.
CITIES = [
    ("Tbilisi", "Tbilisi", "self_governing_city"),
    ("Batumi", "Adjara", "self_governing_city"),
    ("Kobuleti", "Adjara", "statutory_city"),
    ("Lanchkhuti", "Guria", "statutory_city"),
    ("Ozurgeti", "Guria", "statutory_city"),
    ("Kutaisi", "Imereti", "self_governing_city"),
    ("Baghdati", "Imereti", "statutory_city"),
    ("Vani", "Imereti", "statutory_city"),
    ("Zestaponi", "Imereti", "statutory_city"),
    ("Terjola", "Imereti", "statutory_city"),
    ("Samtredia", "Imereti", "statutory_city"),
    ("Sachkhere", "Imereti", "statutory_city"),
    ("Tkibuli", "Imereti", "statutory_city"),
    ("Tskaltubo", "Imereti", "statutory_city"),
    ("Chiatura", "Imereti", "statutory_city"),
    ("Khoni", "Imereti", "statutory_city"),
    ("Akhmeta", "Kakheti", "statutory_city"),
    ("Gurjaani", "Kakheti", "statutory_city"),
    ("Dedoplistskaro", "Kakheti", "statutory_city"),
    ("Telavi", "Kakheti", "statutory_city"),
    ("Lagodekhi", "Kakheti", "statutory_city"),
    ("Sagarejo", "Kakheti", "statutory_city"),
    ("Sighnaghi", "Kakheti", "statutory_city"),
    ("Tsnori", "Kakheti", "statutory_city"),
    ("Kvareli", "Kakheti", "statutory_city"),
    ("Dusheti", "Mtskheta-Mtianeti", "statutory_city"),
    ("Mtskheta", "Mtskheta-Mtianeti", "statutory_city"),
    ("Ambrolauri", "Racha-Lechkhumi and Kvemo Svaneti", "statutory_city"),
    ("Oni", "Racha-Lechkhumi and Kvemo Svaneti", "statutory_city"),
    ("Tsageri", "Racha-Lechkhumi and Kvemo Svaneti", "statutory_city"),
    ("Poti", "Samegrelo-Zemo Svaneti", "self_governing_city"),
    ("Abasha", "Samegrelo-Zemo Svaneti", "statutory_city"),
    ("Zugdidi", "Samegrelo-Zemo Svaneti", "statutory_city"),
    ("Martvili", "Samegrelo-Zemo Svaneti", "statutory_city"),
    ("Senaki", "Samegrelo-Zemo Svaneti", "statutory_city"),
    ("Chkhorotsku", "Samegrelo-Zemo Svaneti", "statutory_city"),
    ("Tsalenjikha", "Samegrelo-Zemo Svaneti", "statutory_city"),
    ("Jvari", "Samegrelo-Zemo Svaneti", "statutory_city"),
    ("Khobi", "Samegrelo-Zemo Svaneti", "statutory_city"),
    ("Akhalkalaki", "Samtskhe-Javakheti", "statutory_city"),
    ("Akhaltsikhe", "Samtskhe-Javakheti", "statutory_city"),
    ("Vale", "Samtskhe-Javakheti", "statutory_city"),
    ("Borjomi", "Samtskhe-Javakheti", "statutory_city"),
    ("Ninotsminda", "Samtskhe-Javakheti", "statutory_city"),
    ("Rustavi", "Kvemo Kartli", "self_governing_city"),
    ("Bolnisi", "Kvemo Kartli", "statutory_city"),
    ("Gardabani", "Kvemo Kartli", "statutory_city"),
    ("Dmanisi", "Kvemo Kartli", "statutory_city"),
    ("Tetritskaro", "Kvemo Kartli", "statutory_city"),
    ("Marneuli", "Kvemo Kartli", "statutory_city"),
    ("Tsalka", "Kvemo Kartli", "statutory_city"),
    ("Gori", "Shida Kartli", "statutory_city"),
    ("Kaspi", "Shida Kartli", "statutory_city"),
    ("Kareli", "Shida Kartli", "statutory_city"),
    ("Khashuri", "Shida Kartli", "statutory_city"),
    ("Gagra", "Abkhazia", "statutory_city_outside_government_control"),
    ("Gali", "Abkhazia", "statutory_city_outside_government_control"),
    ("Gudauta", "Abkhazia", "statutory_city_outside_government_control"),
    ("New Athos", "Abkhazia", "statutory_city_outside_government_control"),
    ("Ochamchire", "Abkhazia", "statutory_city_outside_government_control"),
    ("Sokhumi", "Abkhazia", "statutory_city_outside_government_control"),
    ("Tkvarcheli", "Abkhazia", "statutory_city_outside_government_control"),
    ("Tskhinvali", "Tskhinvali Region", "statutory_city_outside_government_control"),
]

CONTROLLED_CITY_NAMES = [name for name, _, level in CITIES if level != "statutory_city_outside_government_control"]

# Explicit GeoNames IDs avoid same-name villages and transliteration differences.
POINT_OVERRIDES: dict[str, str] = {
    "Kobuleti": "613762",
    "Vani": "611219",
    "Samtredia": "612126",
    "Tkibuli": "611584",
    "Tskaltubo": "824288",
    "Khoni": "613902",
    "Dedoplistskaro": "615005",
    "Oni": "612592",
    "Chkhorotsku": "615140",
    "Jvari": "614740",
    "Akhalkalaki": "615893",
    "Akhaltsikhe": "615860",
    "Ninotsminda": "612691",
    "Rustavi": "612287",
    "Bolnisi": "615419",
    "Dmanisi": "614891",
    "Tetritskaro": "611674",
    "Gori": "614455",
    "New Athos": "612660",
    "Tkvarcheli": "611583",
    "Tskhinvali": "611403",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def require_hash(path: Path, expected: str) -> None:
    actual = sha256(path)
    if actual != expected:
        raise SystemExit(f"{path} 快照哈希变化：{actual}")


def normalized_name(value: str) -> str:
    value = unicodedata.normalize("NFKD", value).casefold()
    return "".join(character for character in value if character.isalnum())


def validate_population_workbook(path: Path) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if len(workbook.sheetnames) != 1:
        raise SystemExit(f"官方人口普查工作簿工作表数量变化：{len(workbook.sheetnames)}")
    sheet = workbook[workbook.sheetnames[0]]
    city_names = [
        value[3:].strip()
        for (value,) in sheet.iter_rows(min_col=1, max_col=1, values_only=True)
        if isinstance(value, str) and value.startswith("C. ")
    ]
    workbook.close()
    if city_names != CONTROLLED_CITY_NAMES:
        missing = sorted(set(CONTROLLED_CITY_NAMES) - set(city_names))
        added = sorted(set(city_names) - set(CONTROLLED_CITY_NAMES))
        raise SystemExit(
            "官方人口普查城市清单变化："
            f"共 {len(city_names)} 个，缺少 {missing}，新增 {added}"
        )


def parse_geonames(
    path: Path,
) -> tuple[dict[str, dict[str, object]], dict[str, list[dict[str, object]]]]:
    by_id: dict[str, dict[str, object]] = {}
    by_name: dict[str, list[dict[str, object]]] = defaultdict(list)
    with path.open(encoding="utf-8") as source:
        for line in source:
            values = line.rstrip("\n").split("\t")
            if len(values) < 19 or values[6] != "P":
                continue
            place: dict[str, object] = {
                "id": values[0],
                "latitude": values[4],
                "longitude": values[5],
                "feature_code": values[7],
                "population": int(values[14] or 0),
            }
            by_id[values[0]] = place
            for name in {values[1], values[2], *values[3].split(",")}:
                key = normalized_name(name)
                if key:
                    by_name[key].append(place)
    return by_id, by_name


def point_rank(place: dict[str, object]) -> tuple[int, int, int]:
    important = int(str(place["feature_code"]) in {"PPLC", "PPLA", "PPLA2", "PPLA3"})
    return int(place["population"]), important, int(place["id"])


def build_rows(
    by_id: dict[str, dict[str, object]],
    by_name: dict[str, list[dict[str, object]]],
) -> list[dict[str, str]]:
    expected_levels = {
        "self_governing_city": 5,
        "statutory_city": 50,
        "statutory_city_outside_government_control": 8,
    }
    if len(CITIES) != 63 or Counter(city[2] for city in CITIES) != expected_levels:
        raise SystemExit("现行法定城市数量或等级构成变化")
    if len({city[0] for city in CITIES}) != 63:
        raise SystemExit("现行法定城市清单存在重名")

    output: list[dict[str, str]] = []
    unresolved: list[str] = []
    for name, admin_area, city_level in CITIES:
        override = POINT_OVERRIDES.get(name)
        if override:
            point = by_id.get(override)
            if point is None:
                raise SystemExit(f"固定中心点不存在：{name} -> {override}")
        else:
            candidates = {
                str(place["id"]): place
                for place in by_name.get(normalized_name(name), [])
            }
            if not candidates:
                unresolved.append(f"{name}\t{admin_area}")
                continue
            point = max(candidates.values(), key=point_rank)

        longitude = float(point["longitude"])
        latitude = float(point["latitude"])
        if not (39.8 <= longitude <= 46.8 and 41.0 <= latitude <= 43.7):
            raise SystemExit(f"{name} 中心点超出格鲁吉亚范围")
        output.append(
            {
                "administrative_code": str(point["id"]),
                "name": name,
                "admin_area": admin_area,
                "city_level": city_level,
                "geonames_id": str(point["id"]),
                "longitude": f"{longitude:.6f}",
                "latitude": f"{latitude:.6f}",
            }
        )

    if unresolved:
        raise SystemExit("城市没有精确中心点：\n" + "\n".join(unresolved))
    for field, values in (
        ("稳定 ID", [row["administrative_code"] for row in output]),
        ("中心坐标", [(row["longitude"], row["latitude"]) for row in output]),
    ):
        duplicates = [value for value, count in Counter(values).items() if count > 1]
        if duplicates:
            raise SystemExit(f"生成结果存在重复{field}：{duplicates}")
    return sorted(output, key=lambda row: (row["admin_area"], row["name"]))


def write_csv(rows: list[dict[str, str]], path: Path) -> None:
    headers = [
        "administrative_code",
        "name",
        "admin_area",
        "city_level",
        "geonames_id",
        "longitude",
        "latitude",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    if len(sys.argv) not in {3, 4}:
        raise SystemExit(
            "用法：py -3.12 scripts/import-georgia-statutory-cities.py "
            "<population-admin-units-2024.xlsx> <GE.txt> [output.csv]"
        )
    population_path = Path(sys.argv[1]).resolve()
    geonames_path = Path(sys.argv[2]).resolve()
    output_path = Path(sys.argv[3]).resolve() if len(sys.argv) == 4 else DEFAULT_OUTPUT
    require_hash(population_path, EXPECTED_POPULATION_HASH)
    require_hash(geonames_path, EXPECTED_GEONAMES_HASH)
    validate_population_workbook(population_path)
    by_id, by_name = parse_geonames(geonames_path)
    rows = build_rows(by_id, by_name)
    write_csv(rows, output_path)
    print("已生成 63 个格鲁吉亚现行法定城市中心点（5 + 50 + 8）")


if __name__ == "__main__":
    main()
