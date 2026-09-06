#!/usr/bin/env python3
"""Build Iran's current county-seat city catalog.

Usage:
  py -3.12 scripts/import-iran-county-seats.py \
    <geo_1404.xlsx> <wikidata-county-capitals.json> <IR.txt> [output.csv]
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


EXPECTED_DIVISIONS_HASH = (
    "4EBF8DE69F64E7634867E096F52346828690DF31975E8638CEC4BF863746C704"
)
EXPECTED_WIKIDATA_HASH = (
    "D684241C6AF4E7526FC24968C5ECC3C17CCA597AF660CCE8F27F1A532952A587"
)
EXPECTED_GEONAMES_HASH = (
    "D950D5EB5C449D7441BD1A3166F7FE46A70ED452199F1E8D7C083D4E7CE669E9"
)

SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_OUTPUT = (
    SCRIPT_DIR.parent / "data" / "sources" / "ir-county-seats-1404.csv"
)

PROVINCE_NAMES = {
    "00": "Markazi",
    "01": "Gilan",
    "02": "Mazandaran",
    "03": "East Azerbaijan",
    "04": "West Azerbaijan",
    "05": "Kermanshah",
    "06": "Khuzestan",
    "07": "Fars",
    "08": "Kerman",
    "09": "Razavi Khorasan",
    "10": "Isfahan",
    "11": "Sistan and Baluchestan",
    "12": "Kurdistan",
    "13": "Hamadan",
    "14": "Chaharmahal and Bakhtiari",
    "15": "Lorestan",
    "16": "Ilam",
    "17": "Kohgiluyeh and Boyer-Ahmad",
    "18": "Bushehr",
    "19": "Zanjan",
    "20": "Semnan",
    "21": "Yazd",
    "22": "Hormozgan",
    "23": "Tehran",
    "24": "Ardabil",
    "25": "Qom",
    "26": "Qazvin",
    "27": "Golestan",
    "28": "North Khorasan",
    "29": "South Khorasan",
    "30": "Alborz",
}

# County code -> display name and GeoNames ID. These counties do not join
# cleanly to the fixed Wikidata snapshot because of newer creations, renamed
# counties, or Persian orthography differences.
MANUAL_CAPITALS = {
    "0103": ("Bandar-e Anzali", "141679"),
    "0104": ("Talesh", "113343"),
    "0227": ("Shirgah", "115001"),
    "0328": ("Leylan", "125710"),
    "0605": ("Bandar-e Mahshahr", "141663"),
    "0630": ("Seydun", "116318"),
    "0731": ("Beyza", "131536"),
    "0733": ("Qaemiyeh", "7451195"),
    "0736": ("Evaz", "135053"),
    "0824": ("Gonbaki", "46023"),
    "0908": ("Sabzevar", "118063"),
    "0913": ("Quchan", "119115"),
    "0916": ("Mashhad", "124665"),
    "0917": ("Nishapur", "122285"),
    "0923": ("Bardaskan", "141349"),
    "0932": ("Torqabeh", "112616"),
    "0942": ("Qadamgah", "120743"),
    "1016": ("Shahin Shahr", "417472"),
    "1102": ("Chabahar", "1161724"),
    "1310": ("Qorveh-e Dargazin", "119160"),
    "1408": ("Saman", "117627"),
    "1409": ("Ben", "140891"),
    "1410": ("Aluni", "143758"),
    "1411": ("Mal-e Khalifeh", "125077"),
    "1412": ("Farrokh Shahr", "120678"),
    "1709": ("Margown", "124771"),
    "2203": ("Bandar Lengeh", "141665"),
    "2213": ("Sardasht", "117116"),
    "2411": ("Aslan Duz", "142703"),
    "2701": ("Bandar Gaz", "141673"),
    "2703": ("Aliabad-e Katul", "144038"),
    "2806": ("Ashkhaneh", "142768"),
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def require_hash(path: Path, expected: str) -> None:
    actual = sha256(path)
    if actual != expected:
        raise SystemExit(f"{path} 快照哈希变化：{actual}")


def normalized_persian(value: str) -> str:
    value = unicodedata.normalize("NFKC", value)
    for old, new in {
        "ي": "ی",
        "ك": "ک",
        "ئ": "ی",
        "ۀ": "ه",
        "ة": "ه",
        "آ": "ا",
        "أ": "ا",
        "إ": "ا",
    }.items():
        value = value.replace(old, new)
    return "".join(character for character in value if character.isalnum())


def read_counties(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if len(workbook.sheetnames) != 1:
        raise SystemExit(f"官方地理工作簿工作表数量变化：{len(workbook.sheetnames)}")
    sheet = workbook[workbook.sheetnames[0]]
    expected_headers = [
        "کد استان",
        "نام استان",
        "کد شهرستان",
        "نام شهرستان",
        "کد بخش",
        "نام بخش",
        "کد دهستان/ شهر",
        "نام دهستان",
        "کد آبادی",
        "نام ",
        "Diag",
        "Coderec",
    ]
    actual_headers = [sheet.cell(1, column).value for column in range(1, 13)]
    if actual_headers != expected_headers:
        raise SystemExit(f"官方地理工作簿列结构变化：{actual_headers}")

    division_counts: Counter[str] = Counter()
    counties: list[dict[str, str]] = []
    for values in sheet.iter_rows(min_row=2, max_col=12, values_only=True):
        division_type = str(values[11]).strip() if values[11] is not None else ""
        if not division_type:
            continue
        division_counts[division_type] += 1
        if division_type != "2":
            continue
        province_code = str(values[0]).strip()
        county_code = str(values[2]).strip()
        counties.append(
            {
                "administrative_code": f"{province_code}{county_code}",
                "province_code": province_code,
                "county_name_fa": str(values[3]).strip(),
            }
        )
    workbook.close()

    expected_counts = {
        "1": 31,
        "2": 484,
        "3": 1193,
        "4": 2777,
        # The raw workbook uses type 5 for both 1,481 cities and 191 urban
        # zones; the source repository separates them during normalization.
        "5": 1672,
        # Two raw settlement record types normalize to 99,317 villages.
        "6": 95415,
        "8": 3902,
    }
    if dict(sorted(division_counts.items())) != expected_counts:
        raise SystemExit(f"官方行政区划数量变化：{dict(division_counts)}")
    if len({county["administrative_code"] for county in counties}) != 484:
        raise SystemExit("官方县代码不唯一")
    if set(PROVINCE_NAMES) != {county["province_code"] for county in counties}:
        raise SystemExit("官方省代码集合变化")
    return counties


def parse_wikidata(path: Path) -> dict[str, dict[str, str]]:
    bindings = json.loads(path.read_text(encoding="utf-8"))["results"]["bindings"]
    if len(bindings) != 475:
        raise SystemExit(f"Wikidata 县治查询结果数量变化：{len(bindings)}")

    by_county: dict[str, dict[str, str]] = {}
    for binding in bindings:
        county_name = binding.get("countyFa", {}).get("value", "")
        if not county_name or "coord" not in binding:
            continue
        key = normalized_persian(county_name.removeprefix("شهرستان"))
        wkt = binding["coord"]["value"]
        match = re.fullmatch(r"Point\(([-0-9.]+) ([-0-9.]+)\)", wkt)
        if not match:
            continue
        value = {
            "name": binding["capitalLabel"]["value"],
            "wikidata_id": binding["capital"]["value"].rsplit("/", 1)[-1],
            "longitude": match.group(1),
            "latitude": match.group(2),
        }
        previous = by_county.get(key)
        if previous is not None and previous != value:
            if (
                previous["wikidata_id"] == value["wikidata_id"]
                and previous["name"] == value["name"]
            ):
                continue
            raise SystemExit(f"Wikidata 县治存在冲突：{county_name}")
        by_county[key] = value
    return by_county


def parse_geonames(path: Path) -> dict[str, dict[str, str]]:
    by_id: dict[str, dict[str, str]] = {}
    with path.open(encoding="utf-8") as source:
        for line in source:
            values = line.rstrip("\n").split("\t")
            if len(values) < 19 or values[6] != "P":
                continue
            by_id[values[0]] = {
                "longitude": values[5],
                "latitude": values[4],
            }
    return by_id


def build_rows(
    counties: list[dict[str, str]],
    wikidata_by_county: dict[str, dict[str, str]],
    geonames_by_id: dict[str, dict[str, str]],
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    wikidata_count = 0
    for county in counties:
        code = county["administrative_code"]
        manual = MANUAL_CAPITALS.get(code)
        if manual:
            name, geonames_id = manual
            point = geonames_by_id.get(geonames_id)
            if point is None:
                raise SystemExit(f"固定 GeoNames 中心点不存在：{code} -> {geonames_id}")
            wikidata_id = ""
        else:
            key = normalized_persian(county["county_name_fa"])
            point = wikidata_by_county.get(key)
            if point is None:
                raise SystemExit(
                    f"县治没有固定中心点：{code} {county['county_name_fa']}"
                )
            name = point["name"]
            geonames_id = ""
            wikidata_id = point["wikidata_id"]
            wikidata_count += 1

        longitude = float(point["longitude"])
        latitude = float(point["latitude"])
        if not (44.0 <= longitude <= 63.5 and 24.0 <= latitude <= 40.5):
            raise SystemExit(f"{name} 中心点超出伊朗范围")
        rows.append(
            {
                "administrative_code": code,
                "name": name,
                "admin_area": PROVINCE_NAMES[county["province_code"]],
                "city_level": "county_seat",
                "geonames_id": geonames_id,
                "wikidata_id": wikidata_id,
                "longitude": f"{longitude:.6f}",
                "latitude": f"{latitude:.6f}",
            }
        )

    if wikidata_count != 452 or len(rows) != 484:
        raise SystemExit(
            f"县治来源数量变化：Wikidata {wikidata_count}，总计 {len(rows)}"
        )
    for field, values in (
        ("稳定 ID", [row["administrative_code"] for row in rows]),
        ("中心坐标", [(row["longitude"], row["latitude"]) for row in rows]),
    ):
        duplicates = [value for value, count in Counter(values).items() if count > 1]
        if duplicates:
            raise SystemExit(f"生成结果存在重复{field}：{duplicates}")
    return sorted(rows, key=lambda row: (row["admin_area"], row["name"], row["administrative_code"]))


def write_csv(rows: list[dict[str, str]], path: Path) -> None:
    headers = [
        "administrative_code",
        "name",
        "admin_area",
        "city_level",
        "geonames_id",
        "wikidata_id",
        "longitude",
        "latitude",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    if len(sys.argv) not in {4, 5}:
        raise SystemExit(
            "用法：py -3.12 scripts/import-iran-county-seats.py "
            "<geo_1404.xlsx> <wikidata-county-capitals.json> <IR.txt> [output.csv]"
        )
    divisions_path = Path(sys.argv[1]).resolve()
    wikidata_path = Path(sys.argv[2]).resolve()
    geonames_path = Path(sys.argv[3]).resolve()
    output_path = Path(sys.argv[4]).resolve() if len(sys.argv) == 5 else DEFAULT_OUTPUT
    require_hash(divisions_path, EXPECTED_DIVISIONS_HASH)
    require_hash(wikidata_path, EXPECTED_WIKIDATA_HASH)
    require_hash(geonames_path, EXPECTED_GEONAMES_HASH)
    rows = build_rows(
        read_counties(divisions_path),
        parse_wikidata(wikidata_path),
        parse_geonames(geonames_path),
    )
    write_csv(rows, output_path)
    print("已生成 484 个伊朗现行县治城市中心点（452 + 32）")


if __name__ == "__main__":
    main()
